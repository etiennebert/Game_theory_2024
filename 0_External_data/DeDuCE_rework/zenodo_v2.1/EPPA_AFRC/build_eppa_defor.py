"""Deforestation attributed to cropland, pasture and forest-plantation expansion by EPPA (AFRC aggregation) region,
from DeDuCE v2.1 (Singh & Persson 2026; Zenodo 10.5281/zenodo.10633818), unamortized attribution (ha of forest loss
in the year, attributed to the land use that replaced it). Output: an Excel workbook whose period tables are formulas
over the annual data (periods editable), plus a long csv."""
import csv, sys, numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

HERE = "/tmp/claude-0/-home-claude/4124e020-d2c1-5af1-8b6d-b33bb2cbec4b/scratchpad/eppa_defor"
OUT = "/mnt/user-data/outputs/deduce_zenodo"
U = "Deforestation attribution, unamortized (ha)"
YEARS = list(range(2001, 2024))
PERIODS = [("2007-2010", 2008, 2010), ("2010-2015", 2011, 2015), ("2015-2020", 2016, 2020), ("2020-2023", 2021, 2023)]
CATS = ["crop", "live", "fors"]
CAT_OF_GROUP = {"Pasture": "live", "Forest plantation": "fors"}   # every other DeDuCE group is a crop

# ------------------------------------------------------------------ EPPA regions (the user's csv; one unquoted comma in REA)
rows = []
with open(f"{HERE}/eppa_afrc_regions_countries.csv", newline="", encoding="utf8") as fh:
    rd = csv.reader(fh); hdr = next(rd)
    for r in rd:
        if not r: continue
        if len(r) == 6: r = [r[0], r[1] + "," + r[2]] + r[3:]
        rows.append(r)
eppa = pd.DataFrame(rows, columns=hdr)
REGIONS = eppa.eppa_region.tolist(); RNAME = dict(zip(eppa.eppa_region, eppa.region_name))
code_region = {}
for _, r in eppa.iterrows():
    for c in [x.strip() for x in r.gtap_regions.split(",")]: code_region[c] = r.eppa_region

# DeDuCE countries without their own GTAP code -> the GTAP residual region of the EPPA file that lists them
RESIDUAL = {"MMR": "XSE", "TLS": "XSE", "PRK": "XEA", "BTN": "XSA", "MDV": "XSA", "TKM": "XSU", "YEM": "XWS",
            "PNG": "XOC", "SLB": "XOC", "FJI": "XOC", "NCL": "XOC", "VUT": "XOC", "FSM": "XOC",
            "LBR": "XWF", "SLE": "XWF", "GNB": "XWF", "GMB": "XWF", "CPV": "XWF", "LBY": "XNF", "LSO": "XSC",
            "SSD": "XEC", "SOM": "XEC", "SYC": "XEC", "ERI": "XEC",
            "BLZ": "XCA", "CUB": "XCB", "BHS": "XCB", "GRD": "XCB", "ATG": "XCB", "LCA": "XCB", "DMA": "XCB", "VCT": "XCB",
            "BRB": "XCB", "KNA": "XCB", "SUR": "XSM", "GUY": "XSM",
            "BIH": "XER", "MKD": "XER", "MNE": "XER", "MDA": "XEE",
            "GUF": "FRA", "REU": "FRA", "GLP": "FRA", "MTQ": "FRA",      # French overseas departments are in France in GTAP
            "SRB and MNE": "SRB",                                          # DeDuCE composite 2001-2005; both parts are ROE
            "SDN and SSD": "XEC"}   # DeDuCE composite 2001-2011; after the split 99 % of the loss is in South Sudan (XEC, AF12)
def eppa_of(iso):
    if iso in code_region: return code_region[iso], f"GTAP code {iso}"
    if iso in RESIDUAL: return code_region[RESIDUAL[iso]], f"via {RESIDUAL[iso]}"
    raise SystemExit(f"no EPPA region for {iso}")

# ------------------------------------------------------------------ DeDuCE
df = pd.read_pickle(f"{HERE}/../deduce/deduce.pkl")
df["cat"] = df["Commodity group"].map(CAT_OF_GROUP).fillna("crop")
countries = df[["ISO", "Producer country", "Continent/Country group"]].drop_duplicates().sort_values("ISO").reset_index(drop=True)
countries[["eppa", "basis"]] = [eppa_of(i) for i in countries.ISO]
cty = df.groupby(["ISO", "cat", "Year"])[U].sum().unstack("Year").reindex(columns=YEARS).fillna(0.0)
cty = cty.reindex(pd.MultiIndex.from_product([countries.ISO, CATS], names=["ISO", "cat"])).fillna(0.0)
assert abs(cty.values.sum() - df[U].sum()) < 1e-6, "hectares lost in the pivot"
reg = cty.join(countries.set_index("ISO")["eppa"], on="ISO").groupby(["eppa", "cat"]).sum().reindex(
    pd.MultiIndex.from_product([REGIONS, CATS], names=["eppa", "cat"])).fillna(0.0)
assert abs(reg.values.sum() - df[U].sum()) < 1e-6

# ------------------------------------------------------------------ workbook
wb = Workbook()
F = "Arial"; fnt = Font(name=F, size=10); bold = Font(name=F, size=10, bold=True); blue = Font(name=F, size=10, color="0000FF")
hdr_fill = PatternFill("solid", fgColor="DDEBF7"); thin = Side(style="thin", color="999999")
def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c); cell.font = bold; cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = Border(bottom=thin)
def set_font(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font == Font(): c.font = fnt
            elif not c.font.name == F: c.font = Font(name=F, size=10, bold=c.font.bold, color=c.font.color)

# README
ws = wb.active; ws.title = "README"
readme = [
 "Deforestation attributed to cropland, pasture and forest-plantation expansion, by EPPA region (AFRC aggregation)",
 "",
 "Source: DeDuCE v2.1 - Singh, C., Persson, U. M., Croft, S., Kastner, T. & West, C. D. (2024). Commodity-driven deforestation, associated carbon emissions and trade 2001-2023 (Version 2.1) [Dataset]. Zenodo. https://doi.org/10.5281/zenodo.10633818",
 "Method: Singh, C. & Persson, U. M. (2026). Global patterns of commodity-driven deforestation and associated carbon emissions. Nature Food 7, 138-151. https://doi.org/10.1038/s43016-026-01305-4",
 "File used: 'DeDuCE_Deforestation attributed to agricultural and forestry commodities (2001-2023).xlsx', column 'Deforestation attribution, unamortized (ha)':",
 "   the forest loss of each year (Hansen/GFC loss, DeDuCE's forest definition) attributed to the land use that replaced it. Gross loss, no regrowth netted out.",
 "   Not the 5-year amortized 'deforestation risk' column, which spreads each year's loss over the following five years of production.",
 "",
 "Categories (sheet country_annual_ha, column category):",
 "   crop = all DeDuCE crop groups (cereals, oilseeds, fruit and nuts, vegetables, roots and tubers, pulses, stimulant and spice crops, fibre crops, sugar crops)",
 "   live = DeDuCE group 'Pasture' (cattle meat and leather)",
 "   fors = DeDuCE group 'Forest plantation'",
 "",
 "Regions: the 28 EPPA regions of eppa_afrc_regions_countries.csv. Countries with their own GTAP code are assigned directly; the 46 DeDuCE countries without one",
 "   go to the EPPA region holding their GTAP residual region (sheet country_mapping, column basis). Two DeDuCE composites: 'SRB and MNE' (2001-2005) -> ROE;",
 "   'SDN and SSD' (2001-2011) -> AF12 (XEC), because 99 % of the post-2012 loss of the pair is in South Sudan (Sudan alone: 10-90 ha/yr).",
 "   French overseas departments (GUF, REU, GLP, MTQ) go with France (EUR), as in GTAP.",
 "",
 "Periods (sheet periods, editable): a period 'A-B' covers the years A+1 to B, so the four periods do not overlap: 2008-2010, 2011-2015, 2016-2020, 2021-2023.",
 "   rates_ha_per_yr = average annual hectares over the period's years; cumulative_ha = the sum over the period's years. Both are formulas over region_annual_ha,",
 "   which is itself SUMIFS over country_annual_ha (the raw data): change a period's first or last year in 'periods' and every table follows.",
 "",
 "Units: hectares. World rows are the sum of the 28 regions and equal DeDuCE's world total.",
 "Built 2026-10-06 from the Zenodo file downloaded 2026-09-18 (folder zenodo_v2.1), script build_eppa_defor.py.",
]
for i, t in enumerate(readme, 1):
    c = ws.cell(row=i, column=1, value=t); c.font = bold if i == 1 else fnt
ws.column_dimensions["A"].width = 160

# periods
wp = wb.create_sheet("periods")
wp.append(["period", "first year", "last year", "years"])
for lab, a, b in PERIODS: wp.append([lab, a, b, None])
for r in range(2, 2 + len(PERIODS)):
    wp.cell(row=r, column=4, value=f"=C{r}-B{r}+1")
    wp.cell(row=r, column=2).font = blue; wp.cell(row=r, column=3).font = blue
wp.cell(row=7, column=1, value="Blue cells are inputs: 'A-B' = years A+1 to B. Edit first/last year and the rate and cumulative sheets recompute.").font = fnt
style_header(wp, 1, 4); wp.column_dimensions["A"].width = 14
NP = len(PERIODS)

# country_annual_ha (raw data)
wc = wb.create_sheet("country_annual_ha")
wc.append(["ISO", "country", "DeDuCE continent group", "EPPA region", "category"] + YEARS)
cinfo = countries.set_index("ISO")
for (iso, cat), vals in cty.iterrows():
    wc.append([iso, cinfo.loc[iso, "Producer country"], cinfo.loc[iso, "Continent/Country group"], cinfo.loc[iso, "eppa"], cat] + [float(v) for v in vals.values])
NC = wc.max_row
style_header(wc, 1, 5 + len(YEARS)); wc.freeze_panes = "F2"
for col in range(6, 6 + len(YEARS)):
    wc.cell(row=1, column=col).number_format = "0"
    for r in range(2, NC + 1): wc.cell(row=r, column=col).number_format = "#,##0"
wc.column_dimensions["B"].width = 26; wc.column_dimensions["C"].width = 24; wc.column_dimensions["D"].width = 11

# region_annual_ha (SUMIFS over the country sheet)
wr = wb.create_sheet("region_annual_ha")
wr.append(["EPPA region", "category", "region name"] + YEARS)
r = 2
for region in REGIONS:
    for cat in CATS:
        wr.cell(row=r, column=1, value=region); wr.cell(row=r, column=2, value=cat); wr.cell(row=r, column=3, value=RNAME[region])
        for j, y in enumerate(YEARS):
            col = L(4 + j); ccol = L(6 + j)
            wr.cell(row=r, column=4 + j, value=f"=SUMIFS(country_annual_ha!{ccol}$2:{ccol}${NC},country_annual_ha!$D$2:$D${NC},$A{r},country_annual_ha!$E$2:$E${NC},$B{r})").number_format = "#,##0"
        r += 1
NR = r - 1
for cat in CATS:   # world rows
    wr.cell(row=r, column=1, value="World"); wr.cell(row=r, column=2, value=cat); wr.cell(row=r, column=3, value="sum of the 28 regions")
    for j in range(len(YEARS)):
        col = L(4 + j)
        wr.cell(row=r, column=4 + j, value=f"=SUMIFS({col}$2:{col}${NR},$B$2:$B${NR},$B{r})").number_format = "#,##0"
    r += 1
style_header(wr, 1, 3 + len(YEARS)); wr.freeze_panes = "D2"
for j in range(len(YEARS)): wr.cell(row=1, column=4 + j).number_format = "0"
wr.column_dimensions["C"].width = 30
YR_HDR = f"region_annual_ha!$D$1:${L(3 + len(YEARS))}$1"
YR_DATA = lambda rows: f"region_annual_ha!$D$2:${L(3 + len(YEARS))}${rows}"

# rates and cumulative (SUMPRODUCT over region_annual_ha with the period bounds)
def period_sheet(name, per_year):
    w = wb.create_sheet(name)
    w.cell(row=1, column=1, value="EPPA region"); w.cell(row=1, column=2, value="region name")
    blocks = CATS + ["total"]
    for bi, cat in enumerate(blocks):
        for pi in range(NP):
            col = 3 + bi * NP + pi
            w.cell(row=1, column=col, value=f'="{cat} "&periods!$A${2 + pi}')
    w.cell(row=2, column=1, value="hectares per year, average over the period's years" if per_year else "hectares, sum over the period's years").font = Font(name=F, size=9, italic=True)
    rr = 3
    for region in REGIONS + ["World"]:
        w.cell(row=rr, column=1, value=region); w.cell(row=rr, column=2, value=RNAME.get(region, "sum of the 28 regions"))
        for bi, cat in enumerate(blocks):
            for pi in range(NP):
                col = 3 + bi * NP + pi; first = f"periods!$B${2 + pi}"; last = f"periods!$C${2 + pi}"
                if cat == "total":
                    cells = "+".join(f"{L(3 + k * NP + pi)}{rr}" for k in range(3))
                    f = f"={cells}"
                elif region == "World":
                    f = f"=SUM({L(col)}3:{L(col)}{rr - 1})"
                else:
                    f = (f"=SUMPRODUCT((region_annual_ha!$A$2:$A${NR}=$A{rr})*(region_annual_ha!$B$2:$B${NR}=\"{cat}\")"
                         f"*({YR_HDR}>={first})*({YR_HDR}<={last})*{YR_DATA(NR)})" + (f"/({last}-{first}+1)" if per_year else ""))
                w.cell(row=rr, column=col, value=f).number_format = "#,##0"
        rr += 1
    style_header(w, 1, 2 + len(blocks) * NP); w.freeze_panes = "C3"
    w.column_dimensions["A"].width = 12; w.column_dimensions["B"].width = 44
    for c in range(3, 3 + len(blocks) * NP): w.column_dimensions[L(c)].width = 15
    w.row_dimensions[1].height = 30
    for c in range(1, 3 + len(blocks) * NP): w.cell(row=rr - 1, column=c).font = bold
    return w
period_sheet("rates_ha_per_yr", True)
period_sheet("cumulative_ha", False)

# country_mapping
wm = wb.create_sheet("country_mapping")
wm.append(["ISO", "country (DeDuCE)", "DeDuCE continent group", "EPPA region", "EPPA region name", "basis", "crop ha 2001-2023", "live ha 2001-2023", "fors ha 2001-2023"])
tot = cty.sum(axis=1).unstack("cat")
for _, c in countries.iterrows():
    wm.append([c.ISO, c["Producer country"], c["Continent/Country group"], c.eppa, RNAME[c.eppa], c.basis] + [float(tot.loc[c.ISO, k]) for k in CATS])
style_header(wm, 1, 9); wm.freeze_panes = "A2"
for col in "GHI":
    for r_ in range(2, wm.max_row + 1): wm[f"{col}{r_}"].number_format = "#,##0"
wm.column_dimensions["B"].width = 30; wm.column_dimensions["C"].width = 24; wm.column_dimensions["E"].width = 44; wm.column_dimensions["F"].width = 16

# order sheets, fonts
wb._sheets = [wb["README"], wb["periods"], wb["rates_ha_per_yr"], wb["cumulative_ha"], wb["region_annual_ha"], wb["country_annual_ha"], wb["country_mapping"]]
for w in wb.worksheets: set_font(w)
xlsx = f"{OUT}/DeDuCE_deforestation_by_EPPA_region_AFRC.xlsx"; wb.save(xlsx); print("saved", xlsx)

# ------------------------------------------------------------------ long csv (values, for scripts) + a pandas check table
long = []
for (region, cat), vals in reg.iterrows():
    for lab, a, b in PERIODS:
        yrs = [y for y in YEARS if a <= y <= b]
        long.append([region, RNAME[region], cat, lab, a, b, len(yrs), vals[yrs].sum(), vals[yrs].mean()])
long = pd.DataFrame(long, columns=["eppa_region", "region_name", "category", "period", "first_year", "last_year", "n_years", "cumulative_ha", "ha_per_yr"])
long.to_csv(f"{OUT}/DeDuCE_deforestation_by_EPPA_region_AFRC_long.csv", index=False)
long.to_pickle(f"{HERE}/long.pkl")
print(long.pivot_table(index="eppa_region", columns=["category", "period"], values="ha_per_yr").reindex(REGIONS).div(1e3).round(1).to_string())
